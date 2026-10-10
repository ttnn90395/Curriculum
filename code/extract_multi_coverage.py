"""
Projet Education -- Extraction multi-cycle de la couverture curriculaire
===========================================================================
Version 3. Corrige les defauts observes sur les vraies donnees (2015):
  - domaines detectes par LETTRE DE CODE (A/B/C), pas par libelle
    d'en-tete (en 2015 les en-tetes de domaine sont generiques)
  - valeurs "manquantes" (7/8/9, hypothese IEA a confirmer) ignorees
    au lieu de faire rejeter la colonne entiere
  - encodage 2019 PROVISOIRE (1->1.0, 2->0.5, 3->0.0): 1 et 3 soutenus
    par un croisement empirique 2019/2023 sur le Grade 8 (n=117 et
    n=22), 2 deduit par elimination (n=5, non valide). A reverifier
    sur le Grade 4 avec validate_numeric_encoding().
  - comparaison inter-cycles sur ITEMS APPARIES (libelles proches),
    car les listes d'items different d'un cycle a l'autre.

STATUT DE TEST: teste par moi sur 2023 G4 (non-regression) et sur
2019/2023 G8 (mecanique). NON teste sur 2015 G4 ni 2019 G4 (fichiers
absents de mon environnement) -- les diagnostics imprimes servent a
verifier ces deux cycles sur ta machine.
"""

from __future__ import annotations
import re
import zipfile
import tempfile
import difflib
import collections
from dataclasses import dataclass, field
from pathlib import Path
import openpyxl
import pandas as pd

ENCODING_2023 = {
    "All or almost all students": 1.0,
    "Only the more advanced students": 0.5,
    "Not included in the curriculum through Grade 4": 0.0,
    "Not included in the curriculum through Grade 8": 0.0,
}
ENCODING_2015 = {
    "Topics Taught to All or Almost All Students": 1.0,
    "Topics Taught to Only the More Able Students (Top Track)": 0.5,
    "Not Included in the Curriculum Through Grade 4": 0.0,
    "Not Included in the Curriculum Through Grade 8": 0.0,
}
ENCODING_2019_PROVISIONAL = {1: 1.0, 2: 0.5, 3: 0.0}

ENCODING_BY_CYCLE = {
    2023: ENCODING_2023,
    2015: ENCODING_2015,
    2019: ENCODING_2019_PROVISIONAL,
}
PROVISIONAL_CYCLES = {2019}

# Hypothese (convention IEA), A CONFIRMER dans la documentation:
# 7 = non applicable, 8 = non administre, 9 = omis/invalide
MISSING_CODES = {7, 8, 9, 97, 98, 99}

G4_MATH_LETTERS = {"A": "Nombre", "B": "Mesure_Geometrie", "C": "Donnees"}
G8_MATH_LETTERS = {"A": "Nombre", "B": "Algebre", "C": "Geometrie", "D": "Donnees"}

MIN_VALID_CELLS = 5
MIN_VALID_SHARE = 0.9


@dataclass
class CycleFile:
    cycle: int
    path: Path
    sheet_name: str
    letters: dict = field(default_factory=lambda: dict(G4_MATH_LETTERS))
    prefix: str = "MA"


@dataclass
class Item:
    cycle: int
    code: str
    domain: str
    label: str
    values: dict  # pays -> valeur numerique encodee (None si manquant)
    n_valid: int
    n_missing: int
    n_other: int
    other_examples: set
    kept: bool


def _open_sheet(cf: CycleFile):
    path = cf.path
    if path.suffix == ".zip":
        tmpdir = Path(tempfile.mkdtemp())
        with zipfile.ZipFile(path) as z:
            xlsx = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if len(xlsx) != 1:
                raise ValueError(f"{path}: attendu 1 xlsx, trouve {xlsx}")
            z.extract(xlsx[0], tmpdir)
            path = tmpdir / xlsx[0]
    return openpyxl.load_workbook(path, data_only=True)[cf.sheet_name]


def read_items(cf: CycleFile) -> list[Item]:
    ws = _open_sheet(cf)
    r1 = [c.value for c in ws[1]]
    r2 = [c.value for c in ws[2]]
    r3 = [c.value for c in ws[3]]
    country_col = r3.index("Country")
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[country_col]]

    pat = re.compile(rf"^{cf.prefix}(\d{{3}})([A-Z])([A-Z])$")

    # Numeros de question dont au moins une colonne a un en-tete "proportion"
    coverage_q = set()
    for j, code in enumerate(r3):
        m = pat.match(str(code)) if code else None
        if m and r1[j] and "proportion" in str(r1[j]).lower():
            coverage_q.add(m.group(1))

    encoding = ENCODING_BY_CYCLE[cf.cycle]
    items = []
    for j, code in enumerate(r3):
        m = pat.match(str(code)) if code else None
        if not m or m.group(1) not in coverage_q or m.group(3) == "T":
            continue
        domain = cf.letters.get(m.group(2))
        if domain is None:
            continue
        label = str(r2[j]).strip() if r2[j] else str(code)

        values, n_valid, n_missing, n_other, other_ex = {}, 0, 0, 0, set()
        for r in rows:
            raw = r[j] if j < len(r) else None
            if raw is None:
                continue
            country = r[country_col]
            if raw in encoding:
                values[country] = encoding[raw]
                n_valid += 1
            elif isinstance(raw, (int, float)) and int(raw) in MISSING_CODES:
                values[country] = None
                n_missing += 1
            else:
                n_other += 1
                if len(other_ex) < 3:
                    other_ex.add(raw)
        share = n_valid / (n_valid + n_other) if (n_valid + n_other) else 0.0
        kept = n_valid >= MIN_VALID_CELLS and share >= MIN_VALID_SHARE
        items.append(Item(cf.cycle, code, domain, label, values,
                          n_valid, n_missing, n_other, other_ex, kept))
    return items


def print_diagnostics(items: list[Item]) -> None:
    print(f"\n  Diagnostic des items ({items[0].cycle if items else '?'}):")
    print(f"  {'code':<9}{'domaine':<18}{'valid':>6}{'manq.':>6}{'autre':>6}  garde  libelle")
    for it in items:
        flag = "oui" if it.kept else "NON"
        extra = f"  autres={sorted(map(str, it.other_examples))}" if it.n_other else ""
        print(f"  {it.code:<9}{it.domain:<18}{it.n_valid:>6}{it.n_missing:>6}{it.n_other:>6}  {flag:<5}  "
              f"{it.label[:40]}{extra}")


def domain_table(items: list[Item]) -> pd.DataFrame:
    rows = []
    for it in items:
        if not it.kept:
            continue
        for country, v in it.values.items():
            if v is not None:
                rows.append(dict(cycle=it.cycle, country=country, domain=it.domain,
                                 code=it.code, value=v))
    df = pd.DataFrame(rows)
    return (df.groupby(["cycle", "country", "domain"])["value"]
              .agg(mean_coverage="mean", n_items="count").reset_index())


def _norm(s: str) -> str:
    return re.sub(r"\W+", " ", re.sub(r"^[a-z]\)\s*", "", s.lower().strip())).strip()


def match_items(a: list[Item], b: list[Item], threshold: float = 0.75):
    """Paires d'items (a_i, b_j) mutuellement les plus proches par
    libelle, similarite >= threshold. Uniquement items retenus."""
    a = [x for x in a if x.kept]; b = [x for x in b if x.kept]
    if not a or not b:
        return []
    sim = lambda x, y: difflib.SequenceMatcher(None, _norm(x.label), _norm(y.label)).ratio()
    best_ab = {x.code: max(b, key=lambda y: sim(x, y)) for x in a}
    best_ba = {y.code: max(a, key=lambda x: sim(x, y)) for y in b}
    pairs = []
    for x in a:
        y = best_ab[x.code]
        if best_ba[y.code].code == x.code and sim(x, y) >= threshold:
            pairs.append((x, y, sim(x, y)))
    return pairs


def validate_numeric_encoding(numeric_cf: CycleFile, text_cf: CycleFile) -> collections.Counter:
    """Croise le code numerique (ex: 2019) avec le texte encode d'un
    autre cycle (2015 ou 2023) sur les sujets apparies et les memes
    pays. Items numeriques lus SANS filtre d'encodage (codes bruts)."""
    ws = _open_sheet(numeric_cf)
    r2 = [c.value for c in ws[2]]
    r3 = [c.value for c in ws[3]]
    cc = r3.index("Country")
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[cc]]
    pat = re.compile(rf"^{numeric_cf.prefix}(\d{{3}})([A-Z])([A-Z])$")

    raw_items = {}
    for j, code in enumerate(r3):
        m = pat.match(str(code)) if code else None
        if m and m.group(3) != "T" and m.group(2) in numeric_cf.letters and r2[j]:
            raw_items[code] = (str(r2[j]).strip(),
                               {r[cc]: r[j] for r in rows if j < len(r) and r[j] is not None})

    text_items = [i for i in read_items(text_cf) if i.kept]
    sim = lambda l1, l2: difflib.SequenceMatcher(None, _norm(l1), _norm(l2)).ratio()
    cross, matched = collections.Counter(), []
    for code, (lab, vals) in raw_items.items():
        best = max(text_items, key=lambda t: sim(lab, t.label), default=None)
        if best is None or sim(lab, best.label) < 0.75:
            continue
        back = max(raw_items, key=lambda c: sim(raw_items[c][0], best.label))
        if back != code:
            continue
        matched.append((lab, best.label))
        for country, v in vals.items():
            tv = best.values.get(country)
            if tv is not None:
                cross[(v, tv)] += 1

    print(f"\n  Validation: codes {numeric_cf.cycle} (numeriques) vs texte {text_cf.cycle}: "
          f"{len(matched)} sujets apparies, {sum(cross.values())} cellules")
    for a, b in matched[:12]:
        print(f"     {a[:48]}  ||  {b[:48]}")
    for code in sorted({c for c, _ in cross}, key=str):
        tot = sum(n for (c, _), n in cross.items() if c == code)
        parts = ", ".join(f"{v}: {sum(n for (c, vv), n in cross.items() if c == code and vv == v) / tot:.0%}"
                          for v in sorted({v for (c, v) in cross if c == code}, reverse=True))
        print(f"     code {code!r} (n={tot}) -> valeur texte 2e cycle: {parts}")
    return cross


def pairwise_trends(items_by_cycle: dict[int, list[Item]], pairs: list[tuple[int, int]],
                    countries: list[str]) -> None:
    """Pour chaque paire de cycles: items apparies (libelles proches),
    puis moyenne par pays sur CES items uniquement, aux deux dates."""
    for c1, c2 in pairs:
        if c1 not in items_by_cycle or c2 not in items_by_cycle:
            continue
        ps = match_items(items_by_cycle[c1], items_by_cycle[c2])
        print(f"\n  {c1} -> {c2}: {len(ps)} items apparies"
              + ("  (2019 PROVISOIRE)" if {c1, c2} & PROVISIONAL_CYCLES else ""))
        for x, y, s in ps:
            print(f"     sim={s:.2f} | {x.label[:50]} || {y.label[:50]}")
        out = []
        for country in countries:
            a, b = [], []
            for x, y, _ in ps:
                vx, vy = x.values.get(country), y.values.get(country)
                if vx is not None and vy is not None:
                    a.append(vx); b.append(vy)
            if a:
                out.append({"pays": country, "n_items": len(a),
                            str(c1): round(sum(a) / len(a), 3),
                            str(c2): round(sum(b) / len(b), 3),
                            "diff": round(sum(b) / len(b) - sum(a) / len(a), 3)})
        if out:
            print(pd.DataFrame(out).to_string(index=False))
        # Detail item par item: T=tous ou presque, A=plus avances, N=non inclus
        sym = {1.0: "T", 0.5: "A", 0.0: "N", None: "?"}
        detail = []
        for x, y, _ in ps:
            row = {"item (cycle 1)": x.label[:46]}
            for country in countries:
                row[country] = f"{sym[x.values.get(country)]}->{sym[y.values.get(country)]}"
            detail.append(row)
        print("  Detail item par item (T=tous, A=plus avances, N=non inclus):")
        print(pd.DataFrame(detail).to_string(index=False))


def grade_grid_check(cf: CycleFile, items: list[Item], limit_grade: int) -> None:
    """Controle INTERNE de l'encodage: croise la valeur de couverture
    avec la grille 'a quels grades le sujet est principalement
    enseigne' du meme item (colonnes <code>1 .. <code>12, marquees
    Yes/Y). Ne depend d'aucun autre cycle.
    limit_grade: 4 pour le Grade 4, 8 pour le Grade 8."""
    ws = _open_sheet(cf)
    r3 = [c.value for c in ws[3]]
    idx = {c: j for j, c in enumerate(r3) if c}
    cc = idx["Country"]
    rows = {r[cc]: r for r in ws.iter_rows(min_row=4, values_only=True) if r[cc]}
    yes = {"yes", "y", "x"}
    seen_vals = collections.Counter()
    cross = collections.defaultdict(collections.Counter)
    n_items = 0
    for it in items:
        if not it.kept or f"{it.code}1" not in idx or f"{it.code}12" not in idx:
            continue
        n_items += 1
        for country, v in it.values.items():
            row = rows.get(country)
            if v is None or row is None:
                continue
            marked = []
            for g in range(1, 13):
                j = idx[f"{it.code}{g}"]
                cell = row[j] if j < len(row) else None
                seen_vals[cell] += 1
                if isinstance(cell, str) and cell.strip().lower() in yes:
                    marked.append(g)
            if not marked:
                cat = "aucun grade renseigne"
            elif min(marked) <= limit_grade:
                cat = f"enseigne a un grade <= {limit_grade}"
            else:
                cat = f"enseigne seulement apres G{limit_grade}"
            cross[v][cat] += 1
    print(f"\n  Controle interne par grille de grades ({cf.cycle}, limite G{limit_grade}): "
          f"{n_items} items avec grille")
    if not n_items:
        print("     aucune grille trouvee -- controle impossible pour ce cycle")
        return
    print(f"     valeurs de cellules rencontrees dans la grille: "
          f"{[(k, n) for k, n in seen_vals.most_common(5)]}")
    names = {1.0: "tous (1.0)", 0.5: "plus avances (0.5)", 0.0: "non inclus (0.0)"}
    for v in (1.0, 0.5, 0.0):
        tot = sum(cross[v].values())
        if not tot:
            continue
        parts = ", ".join(f"{n / tot:.0%} {cat}" for cat, n in cross[v].most_common())
        print(f"     valeur encodee {names[v]} (n={tot}): {parts}")


def dump_labels(items_by_cycle: dict[int, list[Item]], outpath: Path) -> None:
    """Exporte les libelles COMPLETS pour construire un appariement
    manuel et auditable entre cycles (le fuzzy matching ne suffit pas)."""
    outpath.parent.mkdir(parents=True, exist_ok=True)
    rows = [dict(cycle=i.cycle, code=i.code, domain=i.domain, label=i.label, n_valid=i.n_valid)
            for its in items_by_cycle.values() for i in its]
    pd.DataFrame(rows).to_csv(outpath, index=False)
    print(f"\nLibelles complets exportes: {outpath} ({len(rows)} lignes)")


if __name__ == "__main__":
    import sys
    DATA = Path(__file__).resolve().parent.parent / "data" / "raw"
    files = {
        2015: CycleFile(2015, DATA / "timss2015_curriculum_g4.zip", "T15_G4_Mathematics Module"),
        2019: CycleFile(2019, DATA / "timss2019_curriculum_g4.zip", "T19_G4_Mathematics Module"),
        2023: CycleFile(2023, DATA / "timss2023_curriculum_g4.xlsx", "T23_G4_Mathematics Module"),
    }
    items_by_cycle, tables = {}, []
    for year, cf in files.items():
        if not cf.path.exists():
            print(f"MANQUANT: {cf.path}")
            continue
        print(f"\n=== Cycle {year} (G4, maths) ===")
        its = read_items(cf)
        print_diagnostics(its)
        items_by_cycle[year] = its
        tables.append(domain_table(its))

    if "--dump-labels" in sys.argv:
        dump_labels(items_by_cycle, DATA.parent / "processed" / "item_labels_g4_math.csv")

    full = pd.concat(tables, ignore_index=True)
    targets = full[full["country"].isin(["France", "Germany", "Poland"])]
    print("\n" + "=" * 70)
    print("Moyennes par domaine -- NON COMPARABLES entre cycles (listes d'items")
    print("differentes) ; 2019 PROVISOIRE ; a ne pas interpreter comme tendance")
    print("=" * 70)
    print(targets.pivot_table(index=["country", "domain"], columns="cycle",
                              values="mean_coverage").round(3).to_string())
    print("\nItems retenus par cycle (France):")
    print(targets[targets["country"] == "France"].pivot_table(
        index="domain", columns="cycle", values="n_items").to_string())

    print("\n" + "=" * 70)
    print("VALIDATION DE L'ENCODAGE NUMERIQUE 2019 (le plus fiable: 2015, meme libelle)")
    print("=" * 70)
    if 2019 in files and 2015 in items_by_cycle and 2019 in items_by_cycle:
        validate_numeric_encoding(files[2019], files[2015])
    if 2019 in files and 2023 in items_by_cycle and 2019 in items_by_cycle:
        validate_numeric_encoding(files[2019], files[2023])

    print("\n" + "=" * 70)
    print("CONTROLE INTERNE DE L'ENCODAGE (grille des grades, par cycle)")
    print("=" * 70)
    for year, cf in files.items():
        if year in items_by_cycle:
            grade_grid_check(cf, items_by_cycle[year], limit_grade=4)

    print("\n" + "=" * 70)
    print("TENDANCES SUR ITEMS APPARIES, PAR PAIRE DE CYCLES")
    print("=" * 70)
    pairwise_trends(items_by_cycle, [(2015, 2019), (2019, 2023), (2015, 2023)],
                    ["France", "Germany", "Poland"])