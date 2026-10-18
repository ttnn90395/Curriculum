"""Curriculum MIS EN OEUVRE (declare par les enseignants) -- TIMSS G4 maths, 2019 et 2023.
Source: almanacs "Teacher" (xlsx), feuilles ATBM05xx (2019) / ATBM06xx (2023).
  1 = Mostly taught before this year   2 = Mostly taught this year
  3 = Not yet taught (2019: "or just introduced")   + Not applicable / Not administered / Omitted
Colonnes reperees par les LIBELLES de l'en-tete (pas par position fixe).
ATTENTION: la categorie 3 de 2019 inclut "just introduced", pas celle de 2023:
les parts "pas encore enseigne" ne sont pas comparables entre cycles.
Sorties: data/processed/teacher_topics_g4_math.csv + resume France/Allemagne/Pologne.
"""
import re, sys, tempfile, zipfile
from pathlib import Path
import openpyxl
import pandas as pd

RAW = Path("data/raw")
OUT = Path("data/processed/teacher_topics_g4_math.csv")
SOURCES = {
    2019: (RAW / "timss2019_almanacs_g4.zip", "Context/T19/T19_G4_MAT_TeacherAlmanac.xlsx", r"^ATBM05[A-C][A-Z]$"),
    2023: (RAW / "timss2023_almanacs_context_g4.zip", "T23_Almanacs_Context_G4/T23_G4_MAT_TeacherAlmanac.xlsx", r"^ATBM06[A-C][A-Z]$"),
}
FOCUS = ["France", "Germany", "Poland"]


def key_of(label):
    l = label.lower()
    if "before this year" in l:
        return "pct_before"
    if "this year" in l:
        return "pct_this_year"
    if "not yet taught" in l:
        return "pct_not_yet"
    if "not applicable" in l:
        return "pct_na"
    if "not administered" in l:
        return "pct_not_admin"
    if "omitted" in l:
        return "pct_omitted"
    return None


def parse_sheet(ws, cycle, sheet):
    rows = list(ws.iter_rows(values_only=True))
    question = next((str(r[0]) for r in rows[:8] if r and r[0] and str(r[0]).startswith("Question:")), "")
    label = re.sub(r"^Question:.*?mathematics topics\.\s*", "", question).strip()
    hdr_i = next((i for i, r in enumerate(rows[:12])
                  if r and sum(1 for c in r if c and "taught" in str(c).lower()) >= 3), None)
    if hdr_i is None:
        return []
    h1, h2 = rows[hdr_i], rows[hdr_i + 1]
    pct_cols = [j for j, c in enumerate(h2) if c is not None and str(c).strip() == "%"]
    keys = [key_of(str(h1[j])) if h1[j] is not None else None for j in pct_cols]
    if None in keys or len(set(keys)) != len(keys):
        raise ValueError(f"{cycle} {sheet}: en-tetes non reconnus {[h1[j] for j in pct_cols]}")
    out = []
    for r in rows[hdr_i + 2:]:
        if not r or r[0] is None or not isinstance(r[1], (int, float)):
            continue
        d = dict(cycle=cycle, sheet=sheet, topic_label=label, country=str(r[0]).strip(),
                 sample=r[1], valid_n=r[2])
        for j, k in zip(pct_cols, keys):
            v = r[j]
            d[k] = float(v) if isinstance(v, (int, float)) else None
        out.append(d)
    return out


def main():
    rows = []
    for cycle, (zp, member, pat) in SOURCES.items():
        if not zp.exists():
            sys.exit(f"MANQUANT: {zp}")
        tmp = Path(tempfile.mkdtemp())
        zipfile.ZipFile(zp).extract(member, tmp)
        wb = openpyxl.load_workbook(tmp / member, read_only=True, data_only=True)
        names = [s for s in wb.sheetnames if re.match(pat, s)]
        print(f"[{cycle}] {len(names)} feuilles de sujets: {names[0]} .. {names[-1]}")
        for sn in names:
            rows += parse_sheet(wb[sn], cycle, sn)
        wb.close()
    df = pd.DataFrame(rows)
    df["pct_taught_by_now"] = df["pct_before"] + df["pct_this_year"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False)
    print(f"\nEcrit: {OUT} ({len(df)} lignes, {df['country'].nunique()} pays)")
    s = df[["pct_before", "pct_this_year", "pct_not_yet"]].sum(axis=1, min_count=1)
    bad = df[(s - 100).abs() > 1.0].dropna(subset=["pct_before"])
    print(f"Controle somme(cat.1-3) = 100 +-1: {len(bad)} lignes hors tolerance sur {len(df)}")
    pd.set_option("display.width", 250, "display.max_colwidth", 80)
    for cycle in SOURCES:
        sub = df[(df.cycle == cycle) & df.country.isin(FOCUS)]
        print(f"\n===== {cycle}: % surtout enseigne AVANT cette annee / CETTE annee / pas encore =====")
        pv = sub.pivot_table(index=["sheet", "topic_label"], columns="country",
                             values=["pct_before", "pct_this_year", "pct_not_yet"])
        pv = pv.reorder_levels([1, 0], axis=1)
        pv = pv[[(c, k) for c in FOCUS for k in ("pct_before", "pct_this_year", "pct_not_yet") if (c, k) in pv.columns]]
        print(pv.round(0).to_string())
        fr = df[(df.cycle == cycle) & (df.country == "France")]
        print(f"France: part moyenne d'eleves sans reponse enseignant (Omitted+NotAdmin) = "
              f"{(fr['pct_omitted'].fillna(0) + fr['pct_not_admin'].fillna(0)).mean():.1f} %")


if __name__ == "__main__":
    main()
