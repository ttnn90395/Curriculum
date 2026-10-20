#!/usr/bin/env python3
"""fetch_hours_sources.py -- telecharge et INSPECTE (sans interpreter) les sources d'heures d'enseignement.
  1. TIMSS 2019: heures de maths enseignees (implemente), exhibit 12.1-3, grades 4 et 8 (xlsx).
  2. Eurydice 2024/25: resolution du DOI du rapport (temps recommande, 38 systemes) + manuel + questionnaire vierge.
Aucune interpretation: affichage de la structure brute (indices de colonnes) avant tout parsing."""
import re, sys, hashlib, urllib.request
from pathlib import Path
import openpyxl

RAW = Path("data/raw")
RAW.mkdir(parents=True, exist_ok=True)
UA = {"User-Agent": "Mozilla/5.0"}
FILES = {
    "timss2019_instruction_time_m48.xlsx": "https://timss2019.org/reports/wp-content/uploads/download/classroom/12_1-3_instruction-time-M48.xlsx",
    "timss2019_instruction_time_m48.pdf": "https://timss2019.org/reports/wp-content/uploads/download/classroom/12_1-3_instruction-time-M48.pdf",
    "eurydice_instruction_time_2025_manual.pdf": "https://eurydice.eacea.ec.europa.eu/sites/default/files/2025-07/Manual_InstructionTime_2025.pdf",
    "eurydice_instruction_time_2025_questionnaire.xlsx": "https://eurydice.eacea.ec.europa.eu/sites/default/files/2025-07/00_Questionnaire_InstructionTime_2025.xlsx",
}


def fetch(url, timeout=90):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read(), r.headers.get("Content-Type"), r.geturl()


def inspect_xlsx(path, nrows=22, w=14):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"  feuilles: {wb.sheetnames}")
    for sn in wb.sheetnames[:4]:
        ws = wb[sn]
        print(f"  --- feuille '{sn}' ---")
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            cells = [f"{j}={str(v)[:w]}" for j, v in enumerate(r) if v is not None]
            if cells:
                print(f"   L{i:02d}: " + "  ".join(cells)[:230])
            if i >= nrows:
                break
    wb.close()


for name, url in FILES.items():
    print(f"\n== {name}")
    try:
        data, ct, final = fetch(url)
        (RAW / name).write_bytes(data)
        print(f"  {len(data)} octets; {ct}; sha256 {hashlib.sha256(data).hexdigest()[:16]}")
        if name.endswith(".xlsx"):
            inspect_xlsx(RAW / name)
    except Exception as e:
        print(f"  ECHEC: {type(e).__name__}: {e}")

print("\n== Eurydice: resolution du DOI du rapport (10.2797/1556974)")
try:
    data, ct, final = fetch("https://doi.org/10.2797/1556974")
    print(f"  final: {final}\n  type: {ct}; {len(data)} octets")
    if ct and "html" in ct.lower():
        links = sorted(set(re.findall(r'href="([^"]+\.(?:pdf|xlsx|csv|epub)[^"]*)"', data.decode("utf-8", "ignore"), flags=re.I)))
        print(f"  liens de telechargement trouves ({len(links)}):")
        for l in links[:25]:
            print("   ", l)
    else:
        (RAW / "eurydice_instruction_time_2025_report.bin").write_bytes(data)
        print("  contenu non-HTML enregistre: data/raw/eurydice_instruction_time_2025_report.bin")
except Exception as e:
    print(f"  ECHEC: {type(e).__name__}: {e}")
