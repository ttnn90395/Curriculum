"""Explore les almanacs TIMSS "enseignant" pour trouver le curriculum MIS EN OEUVRE
(proportion d'eleves dont l'enseignant declare avoir enseigne chaque sujet de maths).
Etape 1: structure seulement, rien n'est suppose.
Usage: python code/inspect_teacher_almanac.py [--grade 8] [--kw fraction --kw decimal]
"""
import argparse, tempfile, zipfile
from pathlib import Path
import openpyxl

RAW = Path("data/raw")
KW_DEFAULT = ["taught", "topics", "covered", "coverage"]
COUNTRIES = ("France", "Germany", "Poland")


def zips_for(grade):
    pats = {
        2015: [f"timss2015_almanacs_g{grade}.zip"],
        2019: [f"timss2019_almanacs_g{grade}.zip"],
        2023: [f"timss2023_almanacs_context_g{grade}.zip"],
    }
    return {y: [RAW / n for n in ns if (RAW / n).exists()] for y, ns in pats.items()}


def teacher_files(z):
    return [n for n in z.namelist() if "teacher" in n.lower() and n.lower().endswith((".xlsx", ".xlsm"))]


def row_text(row):
    return " | ".join(str(c) for c in row if c is not None)


def explore(year, zpath, kws, max_show=2):
    print(f"\n{'=' * 72}\nCycle {year}: {zpath.name}\n{'=' * 72}")
    z = zipfile.ZipFile(zpath)
    tf = teacher_files(z)
    print("Fichiers enseignant:", tf or "AUCUN")
    if not tf:
        allx = [n for n in z.namelist() if n.lower().endswith((".xlsx", ".xlsm", ".csv", ".pdf", ".html"))]
        print("Autres fichiers (extrait):", allx[:15])
    tmp = Path(tempfile.mkdtemp())
    for name in tf:
        z.extract(name, tmp)
        wb = openpyxl.load_workbook(tmp / name, read_only=True, data_only=True)
        print(f"\n  {name}: {len(wb.sheetnames)} feuilles; debut: {wb.sheetnames[:8]}")
        hits = []
        for sn in wb.sheetnames:
            head = []
            for i, row in enumerate(wb[sn].iter_rows(values_only=True)):
                head.append(row_text(row))
                if i >= 6:
                    break
            if any(k in " ".join(head).lower() for k in kws):
                hits.append((sn, head))
        print(f"  Feuilles dont l'en-tete contient {kws}: {len(hits)}")
        for sn, head in hits[:max_show]:
            print(f"\n  --- feuille '{sn}' (en-tete) ---")
            for h in head[:6]:
                print("   ", h[:230])
            shown = 0
            for row in wb[sn].iter_rows(values_only=True):
                t = row_text(row)
                if any(t.startswith(c) or f"| {c}" in t[:60] for c in COUNTRIES):
                    print("    >>", t[:230])
                    shown += 1
                if shown >= 3:
                    break
        if hits:
            print("\n  Titres des feuilles candidates:", [h[0] for h in hits[:25]])
        wb.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--grade", type=int, default=4, choices=(4, 8))
    ap.add_argument("--kw", action="append")
    a = ap.parse_args()
    kws = [k.lower() for k in (a.kw or KW_DEFAULT)]
    for y, zs in zips_for(a.grade).items():
        if not zs:
            print(f"[{y}] aucun zip d'almanac trouve dans {RAW}")
        for zp in zs:
            explore(y, zp, kws)


if __name__ == "__main__":
    main()
